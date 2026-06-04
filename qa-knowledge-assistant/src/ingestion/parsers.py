"""
Document parsers for different source types.
Each parser extracts text content and metadata from its respective format.
"""

from __future__ import annotations

import logging
import os
from abc import ABC, abstractmethod
from datetime import datetime
from typing import Optional

from src.models import DocumentMetadata, SourceType, DocType

logger = logging.getLogger(__name__)


class DocumentParser(ABC):
    """Base class for document parsers."""

    @abstractmethod
    def parse(self, file_path: str, **kwargs) -> tuple[str, DocumentMetadata]:
        """Parse a document and return (text_content, metadata)."""

    def _detect_doc_type(self, title: str, content: str) -> DocType:
        title_lower = title.lower()
        content_lower = content[:500].lower()

        if any(kw in title_lower for kw in ["requirement", "spec", "prd", "story"]):
            return DocType.REQUIREMENT
        if any(kw in title_lower for kw in ["test", "tc-", "test case", "test plan"]):
            return DocType.TEST_PLAN
        if any(kw in title_lower for kw in ["bug", "defect", "issue"]):
            return DocType.BUG_REPORT
        if any(kw in title_lower for kw in ["design", "architecture", "hld", "lld"]):
            return DocType.DESIGN_DOC
        if any(kw in title_lower for kw in ["api", "endpoint", "swagger"]):
            return DocType.API_DOC
        if "requirement" in content_lower or "shall" in content_lower:
            return DocType.REQUIREMENT
        return DocType.GENERAL


class PDFParser(DocumentParser):
    """Parse PDF documents using PyPDF2."""

    def parse(self, file_path: str, **kwargs) -> tuple[str, DocumentMetadata]:
        try:
            from PyPDF2 import PdfReader
        except ImportError:
            logger.error("PyPDF2 not installed. Install with: pip install PyPDF2")
            raise

        reader = PdfReader(file_path)
        pages = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                pages.append(text.strip())

        content = "\n\n".join(pages)
        title = kwargs.get("title", os.path.basename(file_path))

        info = reader.metadata
        author = None
        created_at = datetime.utcnow()
        if info:
            author = info.get("/Author")
            if info.get("/CreationDate"):
                try:
                    date_str = info["/CreationDate"].replace("D:", "")[:14]
                    created_at = datetime.strptime(date_str, "%Y%m%d%H%M%S")
                except (ValueError, TypeError):
                    pass

        metadata = DocumentMetadata(
            source_type=SourceType.PDF,
            title=title,
            author=author,
            created_at=created_at,
            tags=kwargs.get("tags", []),
            project=kwargs.get("project", "default"),
            doc_type=self._detect_doc_type(title, content),
        )
        return content, metadata


class ExcelParser(DocumentParser):
    """Parse Excel files using openpyxl."""

    def parse(self, file_path: str, **kwargs) -> tuple[str, DocumentMetadata]:
        try:
            from openpyxl import load_workbook
        except ImportError:
            logger.error("openpyxl not installed. Install with: pip install openpyxl")
            raise

        wb = load_workbook(file_path, read_only=True, data_only=True)
        sections = []

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue

            section_lines = [f"## Sheet: {sheet_name}"]
            headers = [str(h) if h else f"Col{i}" for i, h in enumerate(rows[0])]
            section_lines.append("| " + " | ".join(headers) + " |")
            section_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")

            for row in rows[1:]:
                cells = [str(c) if c is not None else "" for c in row]
                section_lines.append("| " + " | ".join(cells) + " |")

            sections.append("\n".join(section_lines))

        wb.close()
        content = "\n\n".join(sections)
        title = kwargs.get("title", os.path.basename(file_path))

        metadata = DocumentMetadata(
            source_type=SourceType.EXCEL,
            title=title,
            tags=kwargs.get("tags", []),
            project=kwargs.get("project", "default"),
            doc_type=self._detect_doc_type(title, content),
        )
        return content, metadata


class MarkdownParser(DocumentParser):
    """Parse Markdown files."""

    def parse(self, file_path: str, **kwargs) -> tuple[str, DocumentMetadata]:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        title = kwargs.get("title")
        if not title:
            for line in content.split("\n"):
                line = line.strip()
                if line.startswith("# "):
                    title = line[2:].strip()
                    break
            else:
                title = os.path.basename(file_path)

        metadata = DocumentMetadata(
            source_type=SourceType.MARKDOWN,
            title=title,
            tags=kwargs.get("tags", []),
            project=kwargs.get("project", "default"),
            doc_type=self._detect_doc_type(title, content),
        )
        return content, metadata


class PlainTextParser(DocumentParser):
    """Parse plain text content (for Confluence, SharePoint, ADO, Email connectors)."""

    def parse(self, file_path: str, **kwargs) -> tuple[str, DocumentMetadata]:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        source_type = kwargs.get("source_type", SourceType.MARKDOWN)
        if isinstance(source_type, str):
            source_type = SourceType(source_type)

        title = kwargs.get("title", os.path.basename(file_path))

        metadata = DocumentMetadata(
            source_type=source_type,
            source_url=kwargs.get("source_url"),
            title=title,
            author=kwargs.get("author"),
            tags=kwargs.get("tags", []),
            project=kwargs.get("project", "default"),
            doc_type=self._detect_doc_type(title, content),
        )
        return content, metadata


# Parser registry
PARSER_REGISTRY: dict[str, type[DocumentParser]] = {
    ".pdf": PDFParser,
    ".xlsx": ExcelParser,
    ".xls": ExcelParser,
    ".md": MarkdownParser,
    ".markdown": MarkdownParser,
    ".txt": PlainTextParser,
}


def get_parser(file_path: str) -> DocumentParser:
    """Get the appropriate parser for a file based on its extension."""
    ext = os.path.splitext(file_path)[1].lower()
    parser_cls = PARSER_REGISTRY.get(ext)
    if not parser_cls:
        logger.warning(f"No parser for extension '{ext}', using PlainTextParser")
        return PlainTextParser()
    return parser_cls()
